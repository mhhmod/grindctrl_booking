from pathlib import Path
from io import BytesIO
import json,zipfile,copy,hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from timing import attach_timing
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

root=Path(__file__).parent
story=json.loads((root/'story.json').read_text())
EMU=9525
def unit(n):return round(n*EMU)
def color(s):return RGBColor.from_string(s.lstrip('#').upper())
P='http://schemas.openxmlformats.org/presentationml/2006/main'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P14='http://schemas.microsoft.com/office/powerpoint/2010/main'

def add_narration(slide,locale):
    """Embed the original MP3 and a standard audio timing node for the eight scene references.
    Native application playback remains untested; HTML/MP4 are the verified masters.
    """
    mp3=root.parent/'grindctrl-part03-voiceover-ar.mp3'
    if not mp3.exists():mp3=root.parent/'input/audio/grindctrl-part03-voiceover-ar.mp3'
    part=Part(PackURI('/ppt/media/grindctrl-part03-voiceover-ar.mp3'),'audio/mpeg',slide.part.package,mp3.read_bytes())
    audio_id=slide.part.relate_to(part,RT.AUDIO);media_id=slide.part.relate_to(part,RT.MEDIA)
    picture=slide.shapes.add_picture(str(root/'assets/icon-mic.png'),unit(148 if locale=='ar' else 1110),unit(681),unit(22),unit(22))
    picture.name='Original Arabic narration — 286 seconds'
    nv=picture._element.find('{'+P+'}nvPicPr');pr=nv.find('{'+P+'}nvPr')
    audio=etree.SubElement(pr,'{'+A+'}audioFile');audio.set('{'+REL+'}link',audio_id)
    extlist=etree.SubElement(pr,'{'+P+'}extLst');ext=etree.SubElement(extlist,'{'+P+'}ext',uri='{DAA4B4D4-6D71-4841-9C94-3DE7FCFB45AD}')
    media=etree.SubElement(ext,'{'+P14+'}media',nsmap={'p14':P14});media.set('{'+REL+'}embed',media_id)
    cpr=nv.find('{'+P+'}cNvPr');click=etree.SubElement(cpr,'{'+A+'}hlinkClick',action='ppaction://media');click.set('{'+REL+'}id','')
    timing=slide._element.find('{'+P+'}timing');clock=timing.find('.//{'+P+'}cTn');children=clock.find('{'+P+'}childTnLst')
    ident=max(int(e.get('id')) for e in timing.iter('{'+P+'}cTn'))+1
    ae=etree.SubElement(children,'{'+P+'}audio',isNarration='1');cm=etree.SubElement(ae,'{'+P+'}cMediaNode',vol='100000',numSld='8',showWhenStopped='1')
    ct=etree.SubElement(cm,'{'+P+'}cTn',id=str(ident),dur=str(round(story['timeline']['duration']*1000)),fill='hold',nodeType='withEffect')
    conds=etree.SubElement(ct,'{'+P+'}stCondLst');etree.SubElement(conds,'{'+P+'}cond',delay='0')
    target=etree.SubElement(cm,'{'+P+'}tgtEl');etree.SubElement(target,'{'+P+'}spTgt',spid=str(picture.shape_id))

sources=copy.deepcopy(story['slides'])
for ap in story['print_appendices']:
 appendix=copy.deepcopy(story['slides'][ap['index']]);appendix['id']='A'+str(len(sources)-7);appendix['static_time']=ap['time'];appendix['title']=ap['title']
 for e in appendix['elements']:
  if e.get('role')=='heading':e['text']=ap['title']
  if e.get('role')=='framing' and e['type']=='text' and ' / 08' in e['text']['en']:e['text']={'en':'Appendix','ar':'ملحق'}
 sources.append(appendix)

for locale in ('en','ar'):
    p=Presentation();p.slide_width=unit(1280);p.slide_height=unit(720)
    p.core_properties.title=story['title'][locale]
    p.core_properties.subject='Customer 360 call experience and source-linked human interaction notes. Proposed workflow; synthetic records.'
    p.core_properties.author='GRINDCTRL'
    for source in sources:
        s=p.slides.add_slide(p.slide_layouts[6])
        targets=[]
        s.background.fill.solid();s.background.fill.fore_color.rgb=color(story['colors']['ink'] if source['dark'] else story['colors']['bg'])
        for e in source['elements']:
            if not (source['static_time']>=e['at'] and (e['until'] is None or source['static_time']<e['until'])):continue
            delay=0
            animated=delay>0 and source['id']!='A1'
            if e['type']=='line':
                x1=1280-e['x'] if locale=='ar' else e['x'];x2=1280-e['x2'] if locale=='ar' else e['x2']
                shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,unit(x1),unit(e['y']),unit(x2),unit(e['y2']))
                shape.line.color.rgb=color(e['color']);shape.line.width=Pt(e['width']*.75)
                if e.get('arrow'):
                    line=shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
                    arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');arrow.set('w','sm');arrow.set('len','sm');line.append(arrow)
                if animated:targets.append((shape.shape_id,delay))
                continue
            x=1280-e['x']-e['w'] if locale=='ar' else e['x']
            if e['type']=='outline':
                shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,unit(x),unit(e['y']),unit(e['w']),unit(e['h']))
                shape.adjustments[0]=.10;shape.fill.background();shape.line.color.rgb=color(e['color']);shape.line.width=Pt(e['width']*.75)
            elif e['type'] in ('rect','ellipse'):
                shape=s.shapes.add_shape(MSO_SHAPE.OVAL if e['type']=='ellipse' else MSO_SHAPE.ROUNDED_RECTANGLE if e['radius'] else MSO_SHAPE.RECTANGLE,unit(x),unit(e['y']),unit(e['w']),unit(e['h']))
                if e.get('radius'):shape.adjustments[0]=min(.12,e['radius']/min(e['w'],e['h']))
                shape.fill.solid();shape.fill.fore_color.rgb=color(e['fill']);shape.line.color.rgb=color(e['line']);shape.line.width=Pt(.75)
            elif e['type']=='image':
                file=root/'assets'/f'{e["key"]}.png'
                iw,ih=Image.open(file).size;scale=min(e['w']/iw,e['h']/ih);w,h=iw*scale,ih*scale
                shape=s.shapes.add_picture(str(file),unit(x+(e['w']-w)/2),unit(e['y']+(e['h']-h)/2),unit(w),unit(h))
            elif e['type']=='text':
                shape=s.shapes.add_textbox(unit(x),unit(e['y']),unit(e['w']),unit(e['h']))
                frame=shape.text_frame;frame.clear();frame.word_wrap=True
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                frame.vertical_anchor=MSO_ANCHOR.TOP
                for i,line in enumerate(e['text'][locale].split('\n')):
                    para=frame.paragraphs[0] if i==0 else frame.add_paragraph()
                    alignment=e['align']
                    if alignment=='center':para.alignment=PP_ALIGN.CENTER
                    elif alignment=='end':para.alignment=PP_ALIGN.LEFT if locale=='ar' else PP_ALIGN.RIGHT
                    else:para.alignment=PP_ALIGN.RIGHT if locale=='ar' else PP_ALIGN.LEFT
                    para.space_before=para.space_after=Pt(0);para.line_spacing=Pt(e['size']*.75*1.12)
                    para._p.get_or_add_pPr().set('rtl','1' if locale=='ar' and not e['ltr'] else '0')
                    run=para.add_run();run.text=line;run.font.size=Pt(e['size']*.75)
                    # Latin IDs/numbers use Manrope; Arabic text uses the explicit CS face.
                    run.font.name='Manrope'
                    run.font.bold=e['bold'];run.font.color.rgb=color(e['color'])
                    run._r.get_or_add_rPr().set('lang','ar-EG' if locale=='ar' else 'en-US')
                    if locale=='ar':
                        script=OxmlElement('a:cs');script.set('typeface','IBM Plex Sans Arabic Medium');run._r.get_or_add_rPr().append(script)
            if animated:targets.append((shape.shape_id,delay))
        # Suppress inherited theme shadows so the workflow stays crisp and flat.
        for sh in s.shapes:
            props=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if props is not None:
                effect=props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
                if effect is None:props.append(OxmlElement('a:effectLst'))
            style=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
            if style is not None:
                effect=style.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
                if effect is not None:effect.set('idx','0')
        attach_timing(s,[],source['motion'],None)
        if source['id']=='S1':add_narration(s,locale)
        notes=[n[locale] for n in source['notes']]
        notes+=['SOURCES (PDF pages):']+[item['file']+' — '+', '.join(map(str,item['pages']))+' — '+item['supports'][locale] for item in story['sources']]
        notes+=['No existing Concrete/Centro relationship. All sample records are fictional. No live connector. Reported workflow has not been independently tested.']
        notes+=['Editable static scene reference. The original continuous MP3 is embedded unchanged on slide 1. Native PowerPoint playback has not been verified; slides advance manually and do not reproduce the full HTML animation choreography. HTML and MP4 are the synchronized audiovisual masters. Twelve pages: eight principal scenes plus four note-entry and exception-state references. Install the bundled Manrope and IBM Plex Sans Arabic fonts for layout fidelity.']
        s.notes_slide.notes_text_frame.text='\n\n'.join(notes)
    dest=root.parent/f'GRINDCTRL_Concrete_Part03_Customer360_{locale}.pptx';p.save(dest)
    # Verify native text and diagrams, speaker notes and shape bounds in the actual archive.
    reopened=Presentation(dest)
    assert len(reopened.slides)==len(sources)
    for i,s in enumerate(reopened.slides):
        assert sum(1 for sh in s.shapes if sh.has_text_frame)>=7
        assert 'fiction' in s.notes_slide.notes_text_frame.text.lower() or 'خيال' in s.notes_slide.notes_text_frame.text
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=p.slide_width+unit(1) and sh.top+sh.height<=p.slide_height+unit(1),(locale,i,sh.left,sh.top)
    with zipfile.ZipFile(dest) as z:
        assert len([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])==len(sources)
        assert hashlib.sha256(z.read('ppt/media/grindctrl-part03-voiceover-ar.mp3')).hexdigest()==story['timeline']['audio_sha256']
    print(locale, f"PowerPoint: {len(sources)} editable light slides (8 scenes + 4 complementary reference states); original MP3 hash verified. Native playback not claimed verified.")
